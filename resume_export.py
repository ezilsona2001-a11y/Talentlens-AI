import io
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from docx.shared import Pt

def build_resume_docx(r):
    doc=Document()
    style=doc.styles["Normal"]
    style.font.name="Arial"
    style.font.size=Pt(10)
    p=doc.add_paragraph()
    run=p.add_run(r.get("name",""))
    run.bold=True
    run.font.size=Pt(22)
    p=doc.add_paragraph(r.get("headline",""))
    p.runs[0].bold=True if p.runs else False
    doc.add_paragraph(" | ".join(x for x in [r.get("email",""),r.get("phone",""),r.get("location","")] if x))
    sections=[
        ("SUMMARY",r.get("summary","")),
        ("SKILLS",", ".join(r.get("skills",[]))),
        ("EDUCATION",r.get("education","")),
        ("EXPERIENCE",r.get("experience","")),
        ("PROJECTS",r.get("projects","")),
        ("CERTIFICATIONS",r.get("certifications","")),
    ]
    for title,text in sections:
        if text:
            p=doc.add_paragraph()
            rr=p.add_run(title)
            rr.bold=True
            doc.add_paragraph(str(text))
    out=io.BytesIO()
    doc.save(out)
    return out.getvalue()

def build_resume_pdf(r):
    out=io.BytesIO()
    doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=40)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="GreenTitle",parent=styles["Title"],textColor=colors.HexColor("#15803d"),fontSize=22))
    styles.add(ParagraphStyle(name="GreenHead",parent=styles["Heading2"],textColor=colors.HexColor("#15803d"),fontSize=12))
    story=[Paragraph(str(r.get("name","")),styles["GreenTitle"])]
    if r.get("headline"): story.append(Paragraph(str(r["headline"]),styles["Normal"]))
    contact=" | ".join(x for x in [r.get("email",""),r.get("phone",""),r.get("location","")] if x)
    if contact: story.append(Paragraph(contact,styles["Normal"]))
    for title,text in [
        ("SUMMARY",r.get("summary","")),
        ("SKILLS",", ".join(r.get("skills",[]))),
        ("EDUCATION",r.get("education","")),
        ("EXPERIENCE",r.get("experience","")),
        ("PROJECTS",r.get("projects","")),
        ("CERTIFICATIONS",r.get("certifications","")),
    ]:
        if text:
            story += [Spacer(1,10),Paragraph(title,styles["GreenHead"]),Paragraph(str(text).replace("&","&amp;"),styles["Normal"])]
    doc.build(story)
    return out.getvalue()
